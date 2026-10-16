"""Disposable XLSX input changes; all model calculations run in LibreOffice."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
import xml.etree.ElementTree as E
import json, subprocess, math, tempfile, shutil, os
import openpyxl
root=Path(__file__).resolve().parents[1]
src=root/'Strategy_Preferred_Coverage_Model.xlsx'
test=Path(tempfile.mkdtemp(prefix='strategy-coverage-check-'))
out=test/'recalculated';out.mkdir(exist_ok=True)
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
E.register_namespace('',ns['s'])
layout=json.loads((root/'model/original-layout.json').read_text())
order=list(layout['sheets'])+['Net NAV','Quarterly path','Native sensitivity','Excel review']
def named(k):
 s,c=layout['names'][k].replace("'",'').replace('$','').split('!');return s,c
cases={'Base':[], 'Severe':[('Dashboard','B4','Severe')], 'ATM-zero':[(*named('ATM_MSTR_Remaining'),0)],'VWAP-low':[('Quarterly path','B12',1),('Quarterly path','B11',94.99)],'VWAP-95':[('Quarterly path','B12',1),('Quarterly path','B11',95)],'VWAP-99':[('Quarterly path','B12',1),('Quarterly path','B11',99)],'VWAP-101':[('Quarterly path','B12',1),('Quarterly path','B11',101)],'BTC-1x':[(*named('BTC_Spot'),24170.77052637783)],'STRK-stock':[(*named('STRK_In_Stock'),'Yes')], 'Gate-off':[('Dashboard','B4','Severe'),(*named('Gate_On'),'No')]}
for name,changes in cases.items():
 with ZipFile(src) as z,ZipFile(test/(name+'.xlsx'),'w',ZIP_DEFLATED) as d:
  for n in z.namelist():
   data=z.read(n)
   relevant=[(c,v)for s,c,v in changes if n==f'xl/worksheets/sheet{order.index(s)+1}.xml']
   if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'):
    tree=E.fromstring(data)
    for cell,val in relevant:
     c=tree.find(f'.//s:c[@r="{cell}"]',ns)
     for x in list(c):c.remove(x)
     if isinstance(val,str):
      c.set('t','inlineStr');E.SubElement(E.SubElement(c,'{'+ns['s']+'}is'),'{'+ns['s']+'}t').text=val
     else:c.set('t','n');E.SubElement(c,'{'+ns['s']+'}v').text=str(val)
    # Source-edit tests require a full rebuild of dependency chains, not saved caches.
    for cell in tree.findall('.//s:c',ns):
     if cell.find('s:f',ns) is not None:
      cached=cell.find('s:v',ns)
      if cached is not None:cell.remove(cached)
    data=E.tostring(tree,encoding='utf-8',xml_declaration=True)
   d.writestr(n,data)
cmd=[os.environ.get('SOFFICE',shutil.which('soffice') or '/tmp/lo-bin/soffice'),'-env:UserInstallation='+(test/'profile').as_uri(),'--headless','--convert-to','xlsx','--outdir',str(out)]+[str(test/(n+'.xlsx'))for n in cases]
subprocess.run(cmd,check=True,capture_output=True,timeout=120)
receipts={}
for name in cases:
 w=openpyxl.load_workbook(out/(name+'.xlsx'),data_only=True)
 errors=[(s.title,c.coordinate,c.value)for s in w for row in s for c in row if c.data_type=='e']
 assert not errors,(name,errors[:5])
 p=w['Quarterly path'];cash=[p.cell(49,c).value for c in range(2,22)];short=p['U51'].value
 assert all(abs(p.cell(57,c).value)<1e-6 for c in range(2,22)),name
 receipts[name]={'errors':len(errors),'endingCash':cash[-1],'unfunded':short,'firstRate':p['B40'].value,'grossCommonTotal':sum(p.cell(46,c).value for c in range(2,22)),'puts':sum(p.cell(44,c).value for c in range(2,22)),'mnav':w['Net NAV']['B20'].value,'btcCover':w['Coverage']['G10'].value,'snapshotYears':w['Coverage']['E10'].value,'snapshotMonths':w['Coverage']['C10'].value,'nativeGateYears':w['Native sensitivity']['F6'].value,'cashResidualMax':max(abs(p.cell(57,c).value) for c in range(2,22))}
assert receipts['Severe']['grossCommonTotal']==0
assert receipts['Severe']['unfunded']>0
assert receipts['ATM-zero']['unfunded']>0
assert receipts['VWAP-low']['firstRate']>receipts['VWAP-95']['firstRate']>.12
assert abs(receipts['VWAP-99']['firstRate']-.12)<1e-10
assert receipts['VWAP-101']['firstRate']<.12
assert receipts['STRK-stock']['snapshotMonths']>receipts['Base']['snapshotMonths']
assert receipts['Gate-off']['grossCommonTotal']>0
assert abs(receipts['Base']['puts']-6713.659)<1e-6
assert abs(receipts['BTC-1x']['btcCover']-1)<1e-5
assert abs(receipts['Base']['nativeGateYears']-4880/1559.673002)<1e-6
(root/'docs/formula-verification.json').write_text(json.dumps({'engine':'LibreOffice 26.8.1.1, disposable copies; not Microsoft Excel','date':'2026-10-16','cases':receipts,'artifactToolLimitation':'Native-table driver changes after import produced #REF in Artifact Tool. Original export and independent LibreOffice changes recalculated without errors. No Microsoft Excel claim.'},indent=2))
print(json.dumps(receipts,indent=2))
